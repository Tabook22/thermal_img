"""Private on-disk document library with page-aware full-text search."""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import sqlite3
import subprocess
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image

TEXT_EDITABLE = {".txt", ".md", ".csv"}
MEDIA = {".mp3", ".wav", ".m4a", ".ogg", ".mp4", ".webm", ".mov"}
SUPPORTED = {".pdf", ".docx", ".xlsx", *TEXT_EDITABLE, ".jpg", ".jpeg", ".png", *MEDIA}
MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_MEDIA_BYTES = 250 * 1024 * 1024
MAX_PAGES = 300


def library_root(storage_root: Path) -> Path:
    root = storage_root / "knowledge"
    root.mkdir(parents=True, exist_ok=True)
    return root


@contextmanager
def connect(storage_root: Path):
    db = sqlite3.connect(library_root(storage_root) / "library.db", timeout=15)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("CREATE TABLE IF NOT EXISTS documents (id TEXT PRIMARY KEY, filename TEXT NOT NULL, stored_name TEXT NOT NULL, sha256 TEXT NOT NULL UNIQUE, uploaded_at TEXT NOT NULL, pages INTEGER NOT NULL, passages INTEGER NOT NULL, wiki_status TEXT NOT NULL DEFAULT 'pending', summary TEXT NOT NULL DEFAULT '')")
    columns = {row["name"] for row in db.execute("PRAGMA table_info(documents)")}
    if "wiki_status" not in columns:
        db.execute("ALTER TABLE documents ADD COLUMN wiki_status TEXT NOT NULL DEFAULT 'pending'")
    if "summary" not in columns:
        db.execute("ALTER TABLE documents ADD COLUMN summary TEXT NOT NULL DEFAULT ''")
    db.execute("CREATE VIRTUAL TABLE IF NOT EXISTS passages_fts USING fts5(text, doc_id UNINDEXED, page UNINDEXED, locator UNINDEXED)")
    try:
        with db:
            yield db
    finally:
        db.close()


def ocr_image(image: Image.Image) -> str:
    executable = shutil.which("tesseract")
    if not executable:
        return ""
    output = io.BytesIO()
    image.convert("RGB").save(output, format="PNG")
    try:
        result = subprocess.run([executable, "stdin", "stdout", "--psm", "6"], input=output.getvalue(), capture_output=True, timeout=30, check=True)
        return result.stdout.decode("utf-8", errors="replace")
    except (OSError, subprocess.SubprocessError):
        return ""


def extract_pages(path: Path) -> list[tuple[int, str, str]]:
    suffix = path.suffix.lower()
    pages: list[tuple[int, str, str]] = []
    if suffix == ".pdf":
        with fitz.open(path) as document:
            if len(document) > MAX_PAGES:
                raise ValueError(f"PDF exceeds the {MAX_PAGES}-page limit")
            for number, page in enumerate(document, 1):
                text = page.get_text("text")
                if len(text.strip()) < 40:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    text = ocr_image(Image.open(io.BytesIO(pix.tobytes("png")))) or text
                pages.append((number, f"Page {number}", text))
    elif suffix == ".docx":
        document = Document(path)
        text = "\n".join([p.text for p in document.paragraphs] + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        pages.append((1, "Document", text))
    elif suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for number, sheet in enumerate(workbook.worksheets[:MAX_PAGES], 1):
                lines = [" | ".join(str(value) for value in row if value is not None) for row in sheet.iter_rows(values_only=True)]
                pages.append((number, f"Sheet {sheet.title}", "\n".join(lines)))
        finally:
            workbook.close()
    elif suffix in MEDIA:
        pages.append((1, "Media", ""))
    elif suffix in {".jpg", ".jpeg", ".png"}:
        with Image.open(path) as image:
            pages.append((1, "Image", ocr_image(image)))
    else:
        pages.append((1, "Document", path.read_text(encoding="utf-8-sig", errors="replace")))
    return pages


def chunks(text: str, size: int = 1100, overlap: int = 160) -> list[str]:
    clean = re.sub(r"[ \t]+", " ", text).strip()
    result = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        if end < len(clean):
            split = clean.rfind(" ", start + size // 2, end)
            if split > start:
                end = split
        segment = clean[start:end].strip()
        if segment:
            result.append(segment)
        if end == len(clean):
            break
        start = max(start + 1, end - overlap)
    return result


def add_document(storage_root: Path, path: Path, filename: str) -> dict:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Unsupported file type")
    if path.stat().st_size > (MAX_MEDIA_BYTES if suffix in MEDIA else MAX_DOCUMENT_BYTES):
        raise ValueError("File exceeds the size limit")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    with connect(storage_root) as db:
        existing = db.execute("SELECT * FROM documents WHERE sha256=?", (digest,)).fetchone()
        if existing:
            return dict(existing)
    pages = extract_pages(path)
    passages = [(text, number, locator) for number, locator, body in pages for text in chunks(body)]
    document_id = uuid.uuid4().hex
    stored_name = document_id + suffix
    destination = library_root(storage_root) / stored_name
    shutil.copyfile(path, destination)
    record = {"id": document_id, "filename": Path(filename).name, "stored_name": stored_name, "sha256": digest,
              "uploaded_at": datetime.now(timezone.utc).isoformat(), "pages": len(pages), "passages": len(passages), "wiki_status": "pending" if passages else "unindexed", "summary": ""}
    try:
        with connect(storage_root) as db:
            db.execute("INSERT INTO documents (id,filename,stored_name,sha256,uploaded_at,pages,passages,wiki_status,summary) VALUES (:id,:filename,:stored_name,:sha256,:uploaded_at,:pages,:passages,:wiki_status,:summary)", record)
            db.executemany("INSERT INTO passages_fts (text,doc_id,page,locator) VALUES (?,?,?,?)", ((text, document_id, number, locator) for text, number, locator in passages))
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    return record


def list_documents(storage_root: Path) -> list[dict]:
    with connect(storage_root) as db:
        return [dict(row) for row in db.execute("SELECT id,filename,uploaded_at,pages,passages,wiki_status,summary FROM documents ORDER BY uploaded_at DESC")]


def get_document(storage_root: Path, document_id: str) -> dict | None:
    with connect(storage_root) as db:
        row = db.execute("SELECT * FROM documents WHERE id=?", (document_id,)).fetchone()
        return dict(row) if row else None


def update_text_document(storage_root: Path, document_id: str, text: str) -> dict:
    record = get_document(storage_root, document_id)
    if not record:
        raise FileNotFoundError("Document not found")
    if Path(record["stored_name"]).suffix.lower() not in TEXT_EDITABLE:
        raise ValueError("Only plain text, Markdown, and CSV files can be edited")
    content = text.encode("utf-8")
    if len(content) > 2 * 1024 * 1024:
        raise ValueError("Text file exceeds the 2 MB editing limit")
    digest = hashlib.sha256(content).hexdigest()
    path = library_root(storage_root) / record["stored_name"]
    original = path.read_bytes()
    passages = chunks(text)
    try:
        with connect(storage_root) as db:
            duplicate = db.execute("SELECT id FROM documents WHERE sha256=? AND id<>?", (digest, document_id)).fetchone()
            if duplicate:
                raise ValueError("Another library file already has this content")
            path.write_bytes(content)
            db.execute("DELETE FROM passages_fts WHERE doc_id=?", (document_id,))
            db.executemany("INSERT INTO passages_fts (text,doc_id,page,locator) VALUES (?,?,?,?)",
                           ((chunk, document_id, 1, "Document") for chunk in passages))
            db.execute("UPDATE documents SET sha256=?,pages=1,passages=?,wiki_status=?,summary='' WHERE id=?",
                       (digest, len(passages), "pending" if passages else "unindexed", document_id))
    except Exception:
        path.write_bytes(original)
        raise
    return get_document(storage_root, document_id)


def delete_document(storage_root: Path, document_id: str) -> bool:
    record = get_document(storage_root, document_id)
    if not record:
        return False
    with connect(storage_root) as db:
        db.execute("DELETE FROM passages_fts WHERE doc_id=?", (document_id,))
        db.execute("DELETE FROM documents WHERE id=?", (document_id,))
    (library_root(storage_root) / record["stored_name"]).unlink(missing_ok=True)
    return True


def search_documents(storage_root: Path, query: str, limit: int = 5) -> list[dict]:
    words = re.findall(r"[^\W_]{2,}", query.lower(), flags=re.UNICODE)
    stop = {"the", "and", "for", "with", "from", "what", "which", "where", "when", "have", "were", "that", "this", "about", "find", "show", "report", "reports"}
    terms = list(dict.fromkeys(word for word in words if word not in stop))[:12]
    if not terms:
        return []
    quoted = ['"' + word.replace('"', '') + '"' for word in terms]
    with connect(storage_root) as db:
        statement = "SELECT f.doc_id,f.page,f.locator,f.text,d.filename,bm25(passages_fts) AS rank FROM passages_fts f JOIN documents d ON d.id=f.doc_id WHERE passages_fts MATCH ? ORDER BY rank LIMIT ?"
        count = max(1, min(limit, 20))
        rows = db.execute(statement, (" AND ".join(quoted), count)).fetchall()
        if not rows:
            rows = db.execute(statement, (" OR ".join(quoted), count)).fetchall()
        return [{"document_id": row["doc_id"], "filename": row["filename"], "page": int(row["page"]), "locator": row["locator"], "excerpt": row["text"][:900]} for row in rows]
